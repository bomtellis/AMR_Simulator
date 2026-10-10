"""Build matching editable Markdown, Word and PDF quick-start guides.

Run with the bundled workspace Python after refreshing images with capture.py.
The two exercise steps and control reference come directly from the in-app help.
"""
from __future__ import annotations

import ast
import html
from pathlib import Path


ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
STEM = HERE / 'AMR-Simulator-Quick-Start'
PICTURES = {
    (0, 1): ('small-map-layout.png', 'The connected small map in Layout'),
    (1, 2): ('ltn-checks.png', 'LTN bay-entry warnings in Project checks'),
}


def content():
    tree = ast.parse((ROOT / 'visualiser' / 'quick_start.py').read_text(encoding='utf-8'))
    values = {}
    for node in tree.body:
        if isinstance(node, ast.Assign) and len(node.targets) == 1:
            name = getattr(node.targets[0], 'id', '')
            if name in {'EXERCISES', 'REFERENCE'}:
                values[name] = ast.literal_eval(node.value)
    if set(values) != {'EXERCISES', 'REFERENCE'} or len(values['EXERCISES']) != 2:
        raise ValueError('The app walkthrough must contain exactly two exercises and a reference.')
    return values['EXERCISES'], values['REFERENCE']


def introduction():
    return ('Use this guide beside the AMR Simulator workbench. Complete the small-map '
            'exercise first, then inspect the supplied LTN project. Help > Guided '
            'walkthrough presents the same steps inside the application. Each exercise '
            'ends with a saved run, investigated result and report export.')


def markdown(exercises, reference):
    lines = ['# AMR Simulator Quick Start Guide', '', introduction(), '',
             '![The two-exercise in-app walkthrough](images/guided-walkthrough.png)', '',
             'The walkthrough stays open while you use the workbench.', '']
    for number, exercise in enumerate(exercises, start=1):
        lines += [f'## Exercise {number} {exercise["title"]}', '']
        for step_no, (title, action, expected, _stage) in enumerate(exercise['steps'], start=1):
            lines += [f'### {step_no} {title}', '', f'**Do this:** {action}', '',
                      f'**Expected state:** {expected}', '']
            picture = PICTURES.get((number - 1, step_no - 1))
            if picture:
                filename, caption = picture
                lines += [f'![{caption}](images/{filename})', '', caption + '.', '']
    lines += ['## Controls reference', '']
    lines += ['![Compact Simple simulation settings](images/compact-simulation-settings.png)', '',
              'Run period groups related controls; Show mode details reveals full diagnostics.', '']
    for title, body in reference:
        lines += [f'### {title}', '', body, '']
    lines += ['## Reading the outcomes', '',
              'The small-map completion count is a reference for the supplied practice input after B-C is connected. '
              'The LTN counts describe a six-hour run of the supplied portable archive. '
              'Different project inputs or software revisions can produce different outcomes; use Review to inspect the run you actually saved.', '']
    return '\n'.join(lines)


def build_docx(exercises, reference):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.66)
    section.left_margin = section.right_margin = Inches(.78)
    styles = doc.styles
    for name, size, bold in [('Normal', 10, False), ('Title', 22, True),
                             ('Heading 1', 16, True), ('Heading 2', 12, True)]:
        style = styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(4 if name == 'Normal' else 6)
    styles['Title'].paragraph_format.space_after = Pt(14)
    for border in styles['Title'].element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    for name in ('Heading 1', 'Heading 2'):
        styles[name].paragraph_format.keep_with_next = True
    doc.add_paragraph('AMR Simulator Quick Start Guide', 'Title')
    doc.add_paragraph(introduction())

    def picture(name, caption, width=6.25):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(HERE / 'images' / name), width=Inches(width))
        c = doc.add_paragraph(caption)
        c.alignment = WD_ALIGN_PARAGRAPH.CENTER
        c.runs[0].font.size = Pt(8.5)
        c.runs[0].font.color.rgb = RGBColor(75, 75, 75)

    picture('guided-walkthrough.png', 'The two-exercise walkthrough stays open beside the workbench.', 5.4)
    for number, exercise in enumerate(exercises, start=1):
        doc.add_heading(f'Exercise {number} {exercise["title"]}', level=1)
        for step_no, (title, action, expected, _stage) in enumerate(exercise['steps'], start=1):
            doc.add_heading(f'{step_no} {title}', level=2)
            p = doc.add_paragraph()
            p.add_run('Do this  ').bold = True
            p.add_run(action)
            p = doc.add_paragraph()
            p.add_run('Expected state  ').bold = True
            p.add_run(expected)
            figure = PICTURES.get((number - 1, step_no - 1))
            if figure:
                picture(*figure, width=4.8 if number == 2 else 5.2)
    doc.add_heading('Controls reference', level=1)
    picture('compact-simulation-settings.png', 'Compact Simple simulation settings with optional mode details.', 5.2)
    for title, body in reference:
        doc.add_heading(title, level=2)
        doc.add_paragraph(body)
    doc.add_heading('Reading the outcomes', level=1)
    doc.add_paragraph('The small-map completion count is a reference for the supplied practice input after B-C is connected. '
                      'The LTN counts describe a six-hour run of the supplied portable archive. Different project inputs or '
                      'software revisions can produce different outcomes; use Review to inspect the run you actually saved.')
    doc.save(STEM.with_suffix('.docx'))


def build_pdf(exercises, reference):
    from PIL import Image as PILImage
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer

    regular = Path('C:/Windows/Fonts/segoeui.ttf')
    bold = Path('C:/Windows/Fonts/segoeuib.ttf')
    if regular.is_file() and bold.is_file():
        pdfmetrics.registerFont(TTFont('SegoeQuick', str(regular)))
        pdfmetrics.registerFont(TTFont('SegoeQuickBold', str(bold)))
        pdfmetrics.registerFontFamily('SegoeQuick', normal='SegoeQuick', bold='SegoeQuickBold')
        family, strong = 'SegoeQuick', 'SegoeQuickBold'
    else:
        family, strong = 'Helvetica', 'Helvetica-Bold'
    sheet = getSampleStyleSheet()
    body = ParagraphStyle('GuideBody', parent=sheet['BodyText'], fontName=family,
                          fontSize=10, leading=14, spaceAfter=6)
    title = ParagraphStyle('GuideTitle', parent=body, fontName=strong, fontSize=22,
                           leading=28, textColor=colors.black, spaceAfter=14)
    h1 = ParagraphStyle('GuideH1', parent=body, fontName=strong, fontSize=16,
                        leading=20, spaceBefore=11, spaceAfter=7, keepWithNext=True)
    h2 = ParagraphStyle('GuideH2', parent=body, fontName=strong, fontSize=12,
                        leading=16, spaceBefore=7, spaceAfter=4, keepWithNext=True)
    caption_style = ParagraphStyle('GuideCaption', parent=body, fontSize=8.5,
                                    leading=11, alignment=TA_CENTER, textColor=colors.HexColor('#4b4b4b'))
    story = [Paragraph('AMR Simulator Quick Start Guide', title), Paragraph(html.escape(introduction()), body)]

    def figure(name, caption, width=6.25):
        path = HERE / 'images' / name
        with PILImage.open(path) as im:
            ratio = im.height / im.width
        return KeepTogether([Image(str(path), width=width * inch, height=width * ratio * inch),
                             Spacer(1, 4), Paragraph(html.escape(caption), caption_style)])

    story.append(figure('guided-walkthrough.png', 'The two-exercise walkthrough stays open beside the workbench.', 5.3))
    for number, exercise in enumerate(exercises, start=1):
        story.append(Paragraph(html.escape(f'Exercise {number} {exercise["title"]}'), h1))
        for step_no, (step_title, action, expected, _stage) in enumerate(exercise['steps'], start=1):
            story.append(Paragraph(html.escape(f'{step_no} {step_title}'), h2))
            story.append(Paragraph('<b>Do this</b>  ' + html.escape(action), body))
            story.append(Paragraph('<b>Expected state</b>  ' + html.escape(expected), body))
            picture = PICTURES.get((number - 1, step_no - 1))
            if picture:
                story.append(figure(*picture, width=4.8 if number == 2 else 5.2))
    story.append(Paragraph('Controls reference', h1))
    story.append(figure('compact-simulation-settings.png', 'Compact Simple simulation settings with optional mode details.', 5.2))
    for section_title, paragraph in reference:
        story.append(Paragraph(html.escape(section_title), h2))
        story.append(Paragraph(html.escape(paragraph), body))
    story.append(Paragraph('Reading the outcomes', h1))
    story.append(Paragraph('The small-map completion count is a reference for the supplied practice input after B-C is connected. '
                           'The LTN counts describe a six-hour run of the supplied portable archive. Different project inputs or '
                           'software revisions can produce different outcomes; use Review to inspect the run you actually saved.', body))

    def footer(canvas, document):
        canvas.setFont(family, 8)
        canvas.setFillColor(colors.HexColor('#555555'))
        canvas.drawString(.76 * inch, .43 * inch, 'AMR Simulator Quick Start Guide')
        canvas.drawRightString(7.74 * inch, .43 * inch, str(document.page))

    pdf = SimpleDocTemplate(str(STEM.with_suffix('.pdf')), pagesize=letter,
                            leftMargin=.78 * inch, rightMargin=.78 * inch,
                            topMargin=.72 * inch, bottomMargin=.72 * inch)
    pdf.build(story, onFirstPage=footer, onLaterPages=footer)


def main():
    exercises, reference = content()
    STEM.with_suffix('.md').write_text(markdown(exercises, reference), encoding='utf-8')
    build_docx(exercises, reference)
    build_pdf(exercises, reference)


if __name__ == '__main__':
    main()
